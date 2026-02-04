import os
import pandas as pd
from pypdf import PdfReader
from typing import List
from docx import Document
from pptx import Presentation

def load_document(file_path: str) -> List[str]:
    ext = os.path.splitext(file_path)[1].lower()

    if ext == ".pdf":
        return load_pdf(file_path)
    elif ext in [".docx"]:
        return load_doc(file_path)
    elif ext == ".txt":
        return load_txt(file_path)
    elif ext == ".csv":
        return load_csv(file_path)
    elif ext in [".xls", ".xlsx"]:
        return load_excel(file_path)
    elif ext in [".pptx", ".ppt"]:
        return load_ppt(file_path)
    else:
        raise ValueError(f"Unsupported file type: {ext}")


def load_pdf(file_path: str) -> List[str]:
    # Load PDF and return list of pages as text
    reader = PdfReader(file_path)
    pages = [page.extract_text() for page in reader.pages]
    return pages

def load_doc(file_path: str) -> List[str]:
    doc = Document(file_path)
    pages: List[str] = []

    current_block = []

    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            # Treat empty line as a block separator
            if current_block:
                pages.append("\n".join(current_block))
                current_block = []
        else:
            current_block.append(text)

    # Flush last block
    if current_block:
        pages.append("\n".join(current_block))

    return pages    

def load_ppt(file_path: str) -> List[str]:
    prs = Presentation(file_path)
    pages: List[str] = []

    for i, slide in enumerate(prs.slides, start=1):
        slide_text = []

        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text = shape.text.strip()
                if text:
                    slide_text.append(text)

        if slide_text:
            page = f"Slide {i}\n" + "\n".join(slide_text)
            pages.append(page)

    return pages

def load_excel(file_path: str) -> List[str]:
    pages: List[str] = []

    sheets = pd.read_excel(file_path, sheet_name=None)

    for sheet_name, df in sheets.items():
        # Drop completely empty rows/columns
        df = df.dropna(how="all").dropna(axis=1, how="all")

        if df.empty:
            continue

        # Convert table to readable text
        text = df.astype(str).apply(
            lambda row: " | ".join(row), axis=1
        ).tolist()

        page_text = f"Sheet: {sheet_name}\n" + "\n".join(text)
        pages.append(page_text)

    return pages

def load_csv(file_path: str) -> List[str]:
    df = pd.read_csv(file_path)

    if df.empty:
        return []

    # Convert rows to readable text
    rows = (
        df.fillna("").astype(str).apply(
            lambda row: " | ".join(row), axis=1
        ).tolist()
    )

    page_text = "\n".join(rows)
    return [page_text]

def load_txt(file_path: str) -> List[str]:
    pages: List[str] = []
    current_block = []

    with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
        for line in f:
            line = line.strip()
            if not line:
                if current_block:
                    pages.append("\n".join(current_block))
                    current_block = []
            else:
                current_block.append(line)

    # Flush last block
    if current_block:
        pages.append("\n".join(current_block))

    return pages


def chunk_text(text: str, chunk_size: int = 500, overlap: int = 50) -> List[str]:
    chunks = []
    start = 0
    while start < len(text):
        end = start + chunk_size
        chunks.append(text[start:end])
        start = end - overlap
    return chunks